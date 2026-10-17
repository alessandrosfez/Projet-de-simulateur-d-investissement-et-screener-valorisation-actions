"""
Tests du databook Excel (databook.py) : le classeur est généré à partir d'un fichier importé (aucun
appel réseau), puis relu avec openpyxl. On vérifie les onglets, que les montants sont des nombres,
qu'une donnée absente s'écrit « n.d. » et jamais 0, et que les saisies Transaction Services de
l'écran se retrouvent dans le classeur.
Usage : pip install -r requirements-dev.txt && pytest -v
"""
import io

import pytest
from openpyxl import load_workbook

import callbacks
import engine as m
from databook import build_databook
from i18n import L

CSV = """# societe = Exemple SA
# devise = EUR
# unite = millions
# secteur = Industrials
poste;2018-12-31;2017-12-31;2016-12-31
revenue;1000;900;800
cost_of_revenue;600;550;500
sga;150;140;130
d_and_a;50;45;40
depreciation;40;36;32
pretax_income;120;100;90
interest_expense;10;9;8
net_income;90;75;65
total_assets;2000;1800;1700
current_assets;700;650;600
current_liabilities;500;480;460
total_equity_incl_minorities;900;820;760
retained_earnings;400;350;300
receivables;200;170;150
inventory;
payables;150;140;130
cash;100;90;80
ppe_net;800;750;700
current_debt;50;60;70
long_term_debt;400;420;440
shares_outstanding;100;100;100
cfo;130;110;95
capex;-60;-55;-50
market_cap;1500;1300;1100
"""


def _context(regime: str | None = None):
    text = CSV if regime is None else CSV.replace("# secteur", f"# regime = {regime}\n# secteur")
    dataset = m.parse_statements_csv(text)["dataset"]
    return callbacks._load_context({"source": "import", "filename": "exemple.csv", "dataset": dataset})


def _workbook(ctx, saved=None, lang="fr"):
    content = build_databook(callbacks.databook_inputs(ctx, saved, lang), lang)
    return load_workbook(io.BytesIO(content))


def _row(ws, label):
    """Première ligne de l'onglet dont la colonne A vaut label, sous forme de liste de valeurs."""
    for row in ws.iter_rows(values_only=True):
        if row[0] == label:
            return list(row)
    raise AssertionError(f"ligne « {label} » absente")


@pytest.mark.parametrize("lang", ["fr", "en"])
def test_databook_has_one_sheet_per_section(lang):
    wb = _workbook(_context(), lang=lang)
    assert wb.sheetnames == [L(lang, f"sheet_{key}") for key in
                             ("summary", "statements", "provenance", "earnings_quality", "working_capital", "scores",
                              "transaction")]


def test_amounts_are_numbers_and_missing_data_is_never_zero():
    ws = _workbook(_context())[L("fr", "sheet_statements")]
    revenue = _row(ws, L("fr", "field_revenue"))
    assert revenue[1] == "revenue"
    assert revenue[2:5] == [1000e6, 900e6, 800e6]
    inventory = _row(ws, L("fr", "field_inventory"))
    assert inventory[2:5] == [L("fr", "na")] * 3
    gross_profit = _row(ws, L("fr", "field_gross_profit"))  # dérivé : CA - coût des ventes
    assert gross_profit[2] == pytest.approx(400e6)


def test_provenance_sheet_mirrors_the_statements_grid():
    ws = _workbook(_context())[L("fr", "sheet_provenance")]
    assert _row(ws, L("fr", "field_revenue"))[2] == L("fr", "prov_imported")
    assert _row(ws, L("fr", "field_inventory"))[2] == L("fr", "prov_missing_not_in_import")


def test_ratios_are_numeric_with_their_calculation_detail():
    ws = _workbook(_context())[L("fr", "sheet_earnings_quality")]
    cfo_rows = [row for row in ws.iter_rows(values_only=True) if row[0] == L("fr", "ratio_cfo_to_net_income")]
    values, detail = cfo_rows[0], cfo_rows[1]
    assert values[1] == pytest.approx(130 / 90)
    assert isinstance(detail[1], str) and "=" in detail[1]


def test_both_altman_models_are_exported():
    ws = _workbook(_context())[L("fr", "sheet_scores")]
    labels = [row[0] for row in ws.iter_rows(values_only=True)]
    assert L("fr", "altman_model_original") in labels and L("fr", "altman_model_z_double_prime") in labels
    z_rows = [row for row in ws.iter_rows(values_only=True) if row[0] == L("fr", "altman_z")]
    # Z d'origine (capitalisation du fichier) puis Z'', chacun sur les trois exercices.
    assert all(isinstance(value, float) for row in z_rows for value in row[1:4])


def test_transaction_inputs_from_the_screen_are_reused():
    ctx = _context()
    saved = {"net_debt_items": [{"label": "Complément de prix", "amount": "25"}],
             "ebitda_adjustments": [{"label": "Restructuration", "amount": "10"}], "pension_tax_rate_pct": 25}
    ws = _workbook(ctx, saved)[L("fr", "sheet_transaction")]
    data_net_debt = 50e6 + 400e6 - 100e6  # dette brute - trésorerie (ni dettes locatives ni retraites)
    assert _row(ws, "Complément de prix")[1] == pytest.approx(25e6)
    assert _row(ws, L("fr", "bridge_line_adjusted_net_debt"))[1] == pytest.approx(data_net_debt + 25e6)
    assert _row(ws, "Restructuration")[1] == pytest.approx(10e6)
    ebitda = (120 + 10 + 50) * 1e6
    assert _row(ws, L("fr", "col_adjusted"))[1] == pytest.approx(ebitda + 10e6)


def test_bank_regime_marks_analyses_not_applicable():
    wb = _workbook(_context(regime="bank"))
    expected = L("fr", "databook_not_applicable", regime=L("fr", "regime_bank"))
    for sheet in ("earnings_quality", "working_capital", "transaction"):
        assert wb[L("fr", f"sheet_{sheet}")]["A3"].value == expected
    # Les états financiers restent exportés.
    assert _row(wb[L("fr", "sheet_statements")], L("fr", "field_revenue"))[2] == 1000e6
