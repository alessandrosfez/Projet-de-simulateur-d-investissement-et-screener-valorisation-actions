"""
Databook Excel de l'analyse : le livrable standard d'une mission de Transaction Services, un onglet
par analyse, avec les chiffres en cellules numériques (réutilisables) et, à côté, le détail du calcul
et la provenance de chaque poste.

Ce module ne calcule rien : il reçoit les résultats du moteur (engine.py), déjà calculés par
callbacks.databook_inputs, et ne fait que les mettre en forme. Ni réseau ni Dash, d'où des tests qui
relisent le classeur produit. Les textes (libellés, formules remplies, provenance) sont ceux de
l'interface (layout.py, i18n.py), pour que l'écran et le classeur disent exactement la même chose.

Règle du reste de l'outil : une donnée absente s'écrit « n.d. » (texte grisé), jamais 0.
"""
import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import layout
from engine import ALTMAN_MODELS, BENEISH_INDICES, BENEISH_REFERENCE_MEANS, FIELD_SCHEMA, PIOTROSKI_CRITERIA, STATEMENTS
from i18n import L

# Formats Excel : montants entiers dans la devise des comptes, ratios selon leur nature.
AMOUNT = "#,##0;-#,##0"
FORMATS = {"amount": AMOUNT, "multiple": '0.00"x"', "percent": "0.0%", "days": "0", "index": "0.000",
           "score": "0.00", "rate": "0.0%"}
SHEETS = ("summary", "statements", "provenance", "earnings_quality", "working_capital", "scores", "transaction")

_TITLE = Font(bold=True, size=14)
_BOLD = Font(bold=True)
_NA = Font(italic=True, color="8A8F99")
_NOTE = Font(italic=True, color="5A5F69")
_HEADER_FILL = PatternFill("solid", fgColor="E6E8EE")
_WRAP = Alignment(wrap_text=True, vertical="top")


class _Sheet:
    """Écriture ligne à ligne d'un onglet : chaque cellule est une valeur brute (None -> « n.d. »), un
    texte, ou un couple (valeur, format) pour un nombre."""

    def __init__(self, workbook: Workbook, title: str, lang: str, first_width: int = 46):
        self.ws = workbook.create_sheet(title[:31])
        self.lang = lang
        self.row = 1
        self.ws.column_dimensions["A"].width = first_width

    def _cell(self, column: int, value, bold: bool = False):
        cell = self.ws.cell(row=self.row, column=column)
        number_format = None
        if isinstance(value, tuple):
            value, number_format = value
        if value is None:
            cell.value, cell.font = L(self.lang, "na"), _NA
            return cell
        cell.value = value
        if number_format:
            cell.number_format = FORMATS.get(number_format, number_format)
        if bold:
            cell.font = _BOLD
        if isinstance(value, str) and "\n" in value:
            cell.alignment = _WRAP
        return cell

    def title(self, text: str):
        self.ws.cell(row=self.row, column=1, value=text).font = _TITLE
        self.row += 2

    def heading(self, text: str):
        self.ws.cell(row=self.row, column=1, value=text).font = Font(bold=True, size=12)
        self.row += 1

    def header(self, values: list):
        for column, value in enumerate(values, start=1):
            cell = self.ws.cell(row=self.row, column=column, value=value)
            cell.font, cell.fill = _BOLD, _HEADER_FILL
        self.row += 1

    def line(self, values: list, bold: bool = False):
        for column, value in enumerate(values, start=1):
            self._cell(column, value, bold=bold)
        self.row += 1

    def note(self, text: str):
        cell = self.ws.cell(row=self.row, column=1, value=text)
        cell.font = _NOTE
        self.row += 1

    def blank(self, n: int = 1):
        self.row += n

    def widths(self, n_columns: int, width: int = 16):
        for column in range(2, n_columns + 2):
            self.ws.column_dimensions[get_column_letter(column)].width = width


def _flat(text: str) -> str:
    """Les infobulles de l'écran séparent leurs paragraphes par une ligne vide (markdown) : une seule
    suffit dans une cellule."""
    return text.replace("\n\n", "\n")


def _years(inputs: dict) -> list[pd.Timestamp]:
    return inputs["normalized"].years


def _year_headers(inputs: dict, lang: str) -> list[str]:
    return [layout.format_year(year, lang) for year in _years(inputs)]


# ---------- Onglets ----------

def _summary(wb: Workbook, inputs: dict, lang: str):
    s = _Sheet(wb, L(lang, "sheet_summary"), lang, first_width=40)
    s.title(L(lang, "databook_title", company=inputs["company"]))
    years_text = ", ".join(_year_headers(inputs, lang))
    regime = inputs["regime"]
    source = (L(lang, "databook_source_import", filename=inputs["identifier"]) if inputs["imported"]
              else L(lang, "databook_source_yahoo", ticker=inputs["identifier"]))
    for label, value in (
        (L(lang, "databook_label_source"), source),
        (L(lang, "label_reporting_currency"), inputs["currency"] or L(lang, "na")),
        (L(lang, "databook_label_years"), years_text),
        (L(lang, "databook_label_regime"), L(lang, f"regime_{regime['regime']}")),
        (L(lang, "databook_label_export_date"), inputs["export_date"]),
    ):
        s.line([label, value])
    s.blank()

    s.heading(L(lang, "databook_key_indicators"))
    s.header([L(lang, "col_metric"), L(lang, "databook_col_period"), L(lang, "databook_col_value"),
              L(lang, "databook_col_reading")])
    latest = _years(inputs)[0] if _years(inputs) else None
    latest_text = layout.format_year(latest, lang) if latest is not None else ""
    not_applicable = L(lang, "databook_not_applicable", regime=L(lang, f"regime_{regime['regime']}"))

    def ratio_row(results_by_year, metric):
        if results_by_year is None:
            return [L(lang, f"ratio_{metric}"), "", not_applicable, ""]
        result = results_by_year[latest][metric]
        return [L(lang, f"ratio_{metric}"), latest_text, (result["value"], layout.RATIO_KINDS[metric]), ""]

    s.line(ratio_row(inputs["earnings"], "cfo_to_net_income"))
    s.line(ratio_row(inputs["earnings"], "sloan_accruals"))
    s.line(ratio_row(inputs["wc_days"], "ccc"))

    if inputs["beneish"]:
        result = inputs["beneish"][0]
        pair = layout.pair_label(result["pair"], lang)
        official, indicative = result["m_score"], result["m_score_indicative"]
        if official["value"] is not None:
            s.line([L(lang, "beneish_m_score"), pair, (official["value"], "score"),
                    L(lang, f"beneish_verdict_{official['verdict']}", pair=pair, value=f"{layout.format_number(official['value'], '.2f', lang)}")])
        elif indicative is not None:
            s.line([L(lang, "beneish_m_score_indicative"), pair, (indicative["value"], "score"),
                    L(lang, "databook_indicative_only")])
        else:
            s.line([L(lang, "beneish_m_score"), pair, None, ""])
    else:
        s.line([L(lang, "beneish_m_score"), "", not_applicable if inputs["beneish"] is None else None, ""])

    if inputs["altman"]:
        model = inputs["altman_default_model"]
        result = inputs["altman"][model][0]
        reading = L(lang, f"zone_{result['zone']}") if result["zone"] else ""
        s.line([f"{L(lang, 'altman_z')} – {L(lang, f'altman_model_{model}')}", layout.format_year(result["year"], lang),
                (result["value"], "score"), reading])
    else:
        s.line([L(lang, "altman_z"), "", not_applicable, ""])

    if inputs["piotroski"]:
        result = inputs["piotroski"][0]
        pair = layout.pair_label((result["year"], result["prior_year"]), lang)
        score = (L(lang, "piotroski_score_text", score=result["score"], n=result["n_available"])
                 if result["score"] is not None else None)
        s.line([L(lang, "piotroski_score"), pair, score, ""])
    else:
        s.line([L(lang, "piotroski_score"), "", not_applicable if inputs["piotroski"] is None else None, ""])

    transaction = inputs["transaction"]
    if transaction is None:
        s.line([L(lang, "section_transaction"), "", not_applicable, ""])
    else:
        period = layout.format_year(transaction["year"], lang)
        s.line([L(lang, "bridge_line_adjusted_net_debt"), period, (transaction["bridge"]["adjusted_net_debt"], "amount"), ""])
        s.line([f"{L(lang, 'effects_ebitda')} – {L(lang, 'col_adjusted')}", period,
                (transaction["ebitda"]["adjusted"], "amount"), ""])
        for key in ("ev_to_ebitda", "net_debt_to_ebitda"):
            for column, figures in (("col_published", transaction["published"]), ("col_adjusted", transaction["adjusted"])):
                s.line([f"{L(lang, f'effects_{key}')} – {L(lang, column)}", period, (figures[key], "multiple"), ""])
    s.blank()
    s.note(L(lang, "scores_caveat"))
    s.note(L(lang, "databook_values_note"))
    s.note(L(lang, "footer_disclaimer"))
    s.ws.column_dimensions["B"].width = 22
    s.ws.column_dimensions["C"].width = 18
    s.ws.column_dimensions["D"].width = 90


def _statements(wb: Workbook, inputs: dict, lang: str, provenance: bool):
    """États normalisés (montants) ou provenance de chaque montant, sur la même grille : clé interne en
    colonne B, celle du modèle d'import, pour retrouver un poste d'un fichier à l'autre."""
    s = _Sheet(wb, L(lang, "sheet_provenance" if provenance else "sheet_statements"), lang)
    normalized = inputs["normalized"]
    s.title(L(lang, "sheet_provenance" if provenance else "section_statements"))
    if not provenance:
        s.note(L(lang, "amounts_unit", currency=inputs["currency"] or L(lang, "na")))
    s.blank()
    for statement in STATEMENTS:
        s.heading(L(lang, f"statement_{statement}"))
        s.header([L(lang, "col_item"), L(lang, "import_col_key"), *_year_headers(inputs, lang)])
        for spec in FIELD_SCHEMA:
            if spec.statement != statement:
                continue
            if provenance:
                cells = [layout.provenance_text(normalized.provenance[spec.key][year], lang) for year in normalized.years]
            else:
                cells = [(normalized.get(spec.key, year), "amount") for year in normalized.years]
            s.line([L(lang, f"field_{spec.key}"), spec.key, *cells])
        s.blank()
    s.ws.column_dimensions["B"].width = 30
    for column in range(3, len(normalized.years) + 3):
        s.ws.column_dimensions[get_column_letter(column)].width = 60 if provenance else 18
    s.ws.freeze_panes = "C2"


def _ratio_sheet(wb: Workbook, inputs: dict, lang: str, title_key: str, results_by_year: dict | None,
                 metrics: tuple[str, ...]):
    s = _Sheet(wb, L(lang, f"sheet_{title_key}"), lang)
    s.title(L(lang, f"section_{title_key}"))
    if results_by_year is None:
        s.line([L(lang, "databook_not_applicable", regime=L(lang, f"regime_{inputs['regime']['regime']}"))])
        return
    years = _years(inputs)
    s.header([L(lang, "col_metric"), *_year_headers(inputs, lang)])
    for metric in metrics:
        kind = layout.RATIO_KINDS[metric]
        s.line([L(lang, f"ratio_{metric}"), *[(results_by_year[year][metric]["value"], kind) for year in years]])
    s.blank()
    s.heading(L(lang, "databook_calculation_detail"))
    s.header([L(lang, "col_metric"), *_year_headers(inputs, lang)])
    for metric in metrics:
        s.line([L(lang, f"ratio_{metric}"),
                *[_flat(layout.ratio_tooltip(results_by_year[year][metric], lang)) for year in years]])
    s.blank()
    for metric in metrics:
        s.note(f"{L(lang, f'ratio_{metric}')} : {L(lang, f'ratio_{metric}_help')}")
    s.widths(len(years), width=34)


def _scores(wb: Workbook, inputs: dict, lang: str):
    s = _Sheet(wb, L(lang, "sheet_scores"), lang)
    s.title(L(lang, "section_scores"))
    s.note(L(lang, "scores_caveat"))
    s.blank()
    not_applicable = L(lang, "databook_not_applicable", regime=L(lang, f"regime_{inputs['regime']['regime']}"))
    widest = 2

    s.heading(L(lang, "section_beneish"))
    if inputs["beneish"] is None:
        s.line([not_applicable])
    elif inputs["beneish"]:
        results = inputs["beneish"]
        s.header([L(lang, "col_metric"), *[layout.pair_label(r["pair"], lang) for r in results],
                  L(lang, "col_reference_manipulators"), L(lang, "col_reference_non_manipulators")])
        for index in BENEISH_INDICES:
            s.line([L(lang, f"beneish_{index}"), *[(r["indices"][index]["value"], "index") for r in results],
                    (BENEISH_REFERENCE_MEANS["manipulators"][index], "index"),
                    (BENEISH_REFERENCE_MEANS["non_manipulators"][index], "index")])
        s.line([L(lang, "beneish_m_score"), *[(r["m_score"]["value"], "score") for r in results]], bold=True)
        if any(r["m_score_indicative"] for r in results):
            s.line([L(lang, "beneish_m_score_indicative"),
                    *[(r["m_score_indicative"]["value"], "score") if r["m_score_indicative"] else "—" for r in results]])
        s.blank()
        s.heading(L(lang, "databook_calculation_detail"))
        for index in BENEISH_INDICES:
            cells = []
            for r in results:
                item = r["indices"][index]
                formula = L(lang, f"formula_beneish_{index}",
                            **{k: layout.format_amount(v, lang) for k, v in item["components"].items()})
                flags = [L(lang, f"flag_{flag}") for flag in item["flags"]]
                cells.append("\n".join([formula, *flags]))
            s.line([L(lang, f"beneish_{index}"), *cells])
        widest = max(widest, len(results) + 2)
    s.blank()

    s.heading(L(lang, "section_altman"))
    if inputs["altman"] is None:
        s.line([not_applicable])
    else:
        s.note(L(lang, "databook_altman_default", model=L(lang, f"altman_model_{inputs['altman_default_model']}")))
        for model in ALTMAN_MODELS:
            results = inputs["altman"][model]
            s.line([L(lang, f"altman_model_{model}")], bold=True)
            s.header([L(lang, "col_metric"), *[layout.format_year(r["year"], lang) for r in results]])
            for x in ALTMAN_MODELS[model]["coefficients"]:
                s.line([L(lang, layout.altman_ratio_label_key(model, x)),
                        *[(r["components"][x]["value"], "index") for r in results]])
            s.line([L(lang, "altman_z"), *[(r["value"], "score") for r in results]], bold=True)
            s.line([L(lang, "altman_zone"), *[L(lang, f"zone_{r['zone']}") if r["zone"] else None for r in results]])
            s.blank()
            widest = max(widest, len(results))

    s.heading(L(lang, "section_piotroski"))
    if inputs["piotroski"] is None:
        s.line([not_applicable])
    elif inputs["piotroski"]:
        results = inputs["piotroski"]
        s.header([L(lang, "col_metric"), *[layout.pair_label((r["year"], r["prior_year"]), lang) for r in results]])
        for criterion in PIOTROSKI_CRITERIA:
            s.line([L(lang, f"piotroski_{criterion}"), *[r["criteria"][criterion]["value"] for r in results]])
        s.line([L(lang, "piotroski_score"),
                *[L(lang, "piotroski_score_text", score=r["score"], n=r["n_available"]) if r["score"] is not None else None
                  for r in results]], bold=True)
    s.widths(widest, width=24)


def _transaction(wb: Workbook, inputs: dict, lang: str):
    s = _Sheet(wb, L(lang, "sheet_transaction"), lang)
    s.title(L(lang, "section_transaction"))
    transaction = inputs["transaction"]
    if transaction is None:
        s.line([L(lang, "databook_not_applicable", regime=L(lang, f"regime_{inputs['regime']['regime']}"))])
        return
    currency = inputs["currency"] or L(lang, "na")
    s.note(L(lang, "transaction_intro", year=layout.format_year(transaction["year"], lang)))
    s.blank()

    s.heading(L(lang, "net_debt_bridge_title"))
    s.header([L(lang, "col_item"), L(lang, "col_amount", currency=currency), L(lang, "col_source")])
    for line in transaction["bridge"]["lines"]:
        label = L(lang, f"bridge_line_{line['key']}", label=line.get("label", ""), rate=f"{layout.format_number(line.get('rate', 0) * 100, 'g', lang)}")
        s.line([label, (line["value"], "amount"), L(lang, f"source_{line['source']}")],
               bold=line["source"] == "subtotal")
    for flag in transaction["bridge"]["flags"]:
        s.note(L(lang, f"flag_{flag}"))
    s.blank()

    s.heading(L(lang, "ebitda_adjustments_title"))
    s.header([L(lang, "col_item"), L(lang, "col_amount", currency=currency), L(lang, "col_source")])
    ebitda = transaction["ebitda"]
    s.line([L(lang, "col_published"), (ebitda["published"], "amount"), L(lang, "source_data")])
    for item in ebitda["adjustments"]:
        s.line([item["label"], (item["amount"], "amount"), L(lang, "source_assumption")])
    s.line([L(lang, "col_adjusted"), (ebitda["adjusted"], "amount"), L(lang, "source_subtotal")], bold=True)
    s.blank()

    s.heading(L(lang, "effects_title"))
    s.header([L(lang, "col_metric"), L(lang, "col_published"), L(lang, "col_adjusted"), L(lang, "col_change")])
    published = {"ebitda": ebitda["published"], "net_debt": transaction["bridge"]["net_debt_data"], **transaction["published"]}
    adjusted = {"ebitda": ebitda["adjusted"], "net_debt": transaction["bridge"]["adjusted_net_debt"], **transaction["adjusted"]}
    for key in layout.EFFECT_ROWS:
        kind = "multiple" if key in ("ev_to_ebitda", "net_debt_to_ebitda") else "amount"
        before, after = published.get(key), adjusted.get(key)
        change = after - before if before is not None and after is not None else None
        s.line([L(lang, f"effects_{key}"), (before, kind), (after, kind), (change, kind)])
    s.blank()
    s.line([L(lang, "label_market_cap"), (transaction["market_cap"], "amount")])
    for note in transaction["notes"]:
        s.note(note)
    s.ws.column_dimensions["B"].width = 22
    s.ws.column_dimensions["C"].width = 22
    s.ws.column_dimensions["D"].width = 22


def build_databook(inputs: dict, lang: str) -> bytes:
    """Classeur .xlsx de l'analyse décrite par inputs (voir callbacks.databook_inputs) : synthèse,
    états normalisés, provenance, qualité des résultats, BFR, scores, dette nette et EBITDA ajusté."""
    wb = Workbook()
    wb.remove(wb.active)
    _summary(wb, inputs, lang)
    _statements(wb, inputs, lang, provenance=False)
    _statements(wb, inputs, lang, provenance=True)
    _ratio_sheet(wb, inputs, lang, "earnings_quality", inputs["earnings"],
                 ("cfo_to_net_income", "fcf_to_net_income", "sloan_accruals"))
    _ratio_sheet(wb, inputs, lang, "working_capital", inputs["wc_days"], ("dso", "dio", "dpo", "ccc"))
    _scores(wb, inputs, lang)
    _transaction(wb, inputs, lang)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
